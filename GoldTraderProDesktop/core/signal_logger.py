"""ثبت سیگنال‌ها، پایش برخورد قیمت به سطوح و خروجی Excel؛ بدون معامله واقعی."""
from pathlib import Path
from datetime import datetime,timezone
import csv,json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
class SignalLogger:
    def __init__(self,root=None):
        self.root=Path(root or Path(__file__).resolve().parents[1]);self.logs=self.root/"logs";self.logs.mkdir(parents=True,exist_ok=True)
        self.json_path=self.logs/"detailed_signals.json";self.csv_path=self.logs/"performance_full.csv"
    def all(self):
        try:return json.loads(self.json_path.read_text(encoding="utf-8"))
        except (OSError,ValueError):return []
    def _save(self,rows):
        tmp=self.json_path.with_suffix(".tmp");tmp.write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding="utf-8");tmp.replace(self.json_path)
        fields=sorted({k for r in rows for k in r if not isinstance(r.get(k),(dict,list))})
        with self.csv_path.open("w",encoding="utf-8-sig",newline="") as f:
            w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows([{k:r.get(k,"") for k in fields} for r in rows])
    def append(self,signal):
        if signal.get("action") not in ("BUY","SELL","WAIT"):raise ValueError("Invalid signal action")
        rows=self.all();row=dict(signal);row["id"]=datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_%f")
        if row["action"] in ("BUY","SELL") and row.get("entry") is not None:
            row["status"]="TRACKING";row["result"]="";row["closed_at"]="";row["exit_price"]=""
        else:row["status"]="WAIT"
        rows.append(row);self._save(rows);return row
    def track_price(self,price):
        """نتیجه سیگنال را با لمس قیمت در تیک‌های واقعی پایش می‌کند؛ PnL واقعی معامله نیست."""
        rows=self.all();changed=False;now=datetime.now(timezone.utc).isoformat();price=float(price)
        for row in rows:
            if row.get("status")!="TRACKING" or row.get("action") not in ("BUY","SELL"):continue
            try:sl=float(row["sl"]);tp1=float(row["tp1"]);tp2=float(row["tp2"])
            except (TypeError,ValueError,KeyError):continue
            action=row["action"];result=None
            if action=="BUY":
                if price<=sl:result="STOP_TOUCHED"
                elif price>=tp2:result="TARGET_2_TOUCHED"
                elif price>=tp1:result="TARGET_1_TOUCHED"
            else:
                if price>=sl:result="STOP_TOUCHED"
                elif price<=tp2:result="TARGET_2_TOUCHED"
                elif price<=tp1:result="TARGET_1_TOUCHED"
            if result:
                row["status"]="CLOSED_BY_PRICE_TOUCH";row["result"]=result;row["closed_at"]=now;row["exit_price"]=price;row["price_move_dollars"]=round((price-float(row["entry"])) if action=="BUY" else (float(row["entry"])-price),2);changed=True
        if changed:self._save(rows)
        return changed
    def export_excel(self,path=None):
        rows=self.all();path=Path(path or self.logs/"signals_report.xlsx")
        wb=Workbook();ws=wb.active;ws.title="سیگنال‌ها"
        fields=["id","issued_at","action","entry","sl","tp1","tp2","score","risk","status","result","exit_price","price_move_dollars"]
        ws.append(fields)
        for row in rows:ws.append([row.get(k,"") for k in fields])
        for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="243247")
        detail=wb.create_sheet("دلایل");detail.append(["شناسه","نوع","دلایل"])
        for row in rows:detail.append([row.get("id",""),row.get("action","")," | ".join(row.get("reasons",[])) if isinstance(row.get("reasons"),list) else row.get("reason","")])
        stats=wb.create_sheet("آمار");closed=[r for r in rows if r.get("result") in ("TARGET_1_TOUCHED","TARGET_2_TOUCHED","STOP_TOUCHED")];targets=sum(r.get("result") in ("TARGET_1_TOUCHED","TARGET_2_TOUCHED") for r in closed);stats.append(["شاخص","مقدار"]);stats.append(["کل سیگنال‌ها",len(rows)]);stats.append(["خرید",sum(r.get("action")=="BUY" for r in rows)]);stats.append(["فروش",sum(r.get("action")=="SELL" for r in rows)]);stats.append(["انتظار",sum(r.get("action")=="WAIT" for r in rows)]);stats.append(["هدف اول لمس شد",sum(r.get("result")=="TARGET_1_TOUCHED" for r in rows)]);stats.append(["هدف دوم لمس شد",sum(r.get("result")=="TARGET_2_TOUCHED" for r in rows)]);stats.append(["حد ضرر لمس شد",sum(r.get("result")=="STOP_TOUCHED" for r in rows)]);stats.append(["نرخ سیگنال‌های رسیده به هدف (%)",round(100*targets/len(closed),2) if closed else "NO DATA"]);stats.append(["مجموع حرکت جهت‌دار قیمت ($)",round(sum(float(r.get("price_move_dollars",0) or 0) for r in closed),2) if closed else "NO DATA"]);stats.append(["یادداشت","این آمار لمس قیمت است، نه سود/زیان حساب معاملاتی واقعی."])
        for sh in wb.worksheets:
            sh.freeze_panes="A2";sh.auto_filter.ref=sh.dimensions
            for col in sh.columns:
                letter=col[0].column_letter;sh.column_dimensions[letter].width=min(52,max(14,max(len(str(c.value or "")) for c in col)+2))
        wb.save(path);return path
